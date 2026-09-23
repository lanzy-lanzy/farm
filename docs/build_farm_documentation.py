from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "farm_web_app_documentation.docx"
ARCHITECTURE = DOCS / "generated_architecture.png"
MODULE_MAP = DOCS / "generated_module_map.png"
ACTIVITY = DOCS / "activity_diagram.png"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 28, 36)
MUTED = RGBColor(90, 96, 105)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, total_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_paragraph(doc, text="", style=None, before=0, after=6, line=1.10, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        font(run, size=11, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.text = ""
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    run = p.add_run(text)
    font(run, size=16 if level == 1 else 13 if level == 2 else 12, color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_bullet(doc, text):
    p = add_paragraph(doc, style="List Bullet", after=4, line=1.167)
    p.add_run(text)
    for run in p.runs:
        font(run, size=11, color=INK)
    return p


def add_numbered(doc, text):
    p = add_paragraph(doc, style="List Number", after=4, line=1.167)
    p.add_run(text)
    for run in p.runs:
        font(run, size=11, color=INK)
    return p


def add_table(doc, headers, rows, widths_dxa, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa, sum(widths_dxa), 120)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for idx, label in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        font(run, size=font_size, color=INK, bold=True)
    for row in rows:
        row_obj = table.add_row()
        set_row_cant_split(row_obj)
        cells = row_obj.cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            run = p.add_run(str(value))
            font(run, size=font_size, color=INK)
            if len(str(value)) < 18 and idx != 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths_dxa, sum(widths_dxa), 120)
    add_paragraph(doc, after=4)
    return table


def add_caption(doc, text):
    p = add_paragraph(doc, text, before=2, after=8, line=1.10, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        font(run, size=9.5, color=MUTED, italic=True)
    return p


def draw_box(draw, box, text, fill, outline="#2E74B5", font_obj=None):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    lines = []
    words = text.split()
    line = ""
    for word in words:
        candidate = (line + " " + word).strip()
        if draw.textlength(candidate, font=font_obj) > (x2 - x1 - 40) and line:
            lines.append(line)
            line = word
        else:
            line = candidate
    lines.append(line)
    total_h = len(lines) * 28
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        w = draw.textlength(line, font=font_obj)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, fill="#111827", font=font_obj)
        y += 28


def draw_arrow(draw, start, end):
    draw.line([start, end], fill="#334155", width=4)
    ex, ey = end
    sx, sy = start
    if ex > sx:
        pts = [(ex, ey), (ex - 16, ey - 9), (ex - 16, ey + 9)]
    elif ex < sx:
        pts = [(ex, ey), (ex + 16, ey - 9), (ex + 16, ey + 9)]
    elif ey > sy:
        pts = [(ex, ey), (ex - 9, ey - 16), (ex + 9, ey - 16)]
    else:
        pts = [(ex, ey), (ex - 9, ey + 16), (ex + 9, ey + 16)]
    draw.polygon(pts, fill="#334155")


def generate_diagrams():
    try:
        title_font = ImageFont.truetype("arial.ttf", 30)
        box_font = ImageFont.truetype("arial.ttf", 22)
        small_font = ImageFont.truetype("arial.ttf", 18)
    except OSError:
        title_font = ImageFont.load_default()
        box_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    img = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((420, 35), "System Architecture", fill="#0F172A", font=title_font)
    boxes = {
        "users": (70, 190, 350, 330),
        "templates": (500, 110, 820, 250),
        "django": (500, 340, 820, 500),
        "db": (500, 610, 820, 750),
        "static": (1030, 110, 1370, 250),
        "reports": (1030, 340, 1370, 500),
        "alerts": (1030, 610, 1370, 750),
    }
    draw_box(draw, boxes["users"], "Farm owner, admin, and caretaker", "#F8FAFC", font_obj=box_font)
    draw_box(draw, boxes["templates"], "Django templates with Tailwind CSS", "#E8F2EE", font_obj=box_font)
    draw_box(draw, boxes["django"], "Django views, forms, auth, and models", "#E8EEF5", font_obj=box_font)
    draw_box(draw, boxes["db"], "SQLite development database", "#F2F4F7", font_obj=box_font)
    draw_box(draw, boxes["static"], "Local static assets: HTMX, Alpine, Chart.js", "#FFF7ED", font_obj=box_font)
    draw_box(draw, boxes["reports"], "Reports and Excel export", "#F0F9FF", font_obj=box_font)
    draw_box(draw, boxes["alerts"], "Notifications and activity logs", "#FEF2F2", font_obj=box_font)
    draw_arrow(draw, (350, 260), (500, 180))
    draw_arrow(draw, (660, 250), (660, 340))
    draw_arrow(draw, (660, 500), (660, 610))
    draw_arrow(draw, (820, 180), (1030, 180))
    draw_arrow(draw, (820, 420), (1030, 420))
    draw_arrow(draw, (820, 690), (1030, 690))
    draw.text((75, 805), "The system keeps form submission, records, dashboards, reports, and alerts inside one Django project.", fill="#475569", font=small_font)
    img.save(ARCHITECTURE)

    img = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((505, 35), "Module Map", fill="#0F172A", font=title_font)
    center = (600, 360, 1000, 500)
    draw_box(draw, center, "Tambulig Poultry Farm Web App", "#E8EEF5", font_obj=box_font)
    modules = [
        ("Dashboard", (90, 110, 360, 230)),
        ("Flocks", (90, 350, 360, 470)),
        ("Inventory", (90, 590, 360, 710)),
        ("Feeding", (510, 110, 780, 230)),
        ("Medicine", (820, 110, 1090, 230)),
        ("Eggs", (1210, 110, 1480, 230)),
        ("Mortality", (1210, 350, 1480, 470)),
        ("Sales and Expenses", (1210, 590, 1480, 710)),
        ("Reports", (820, 590, 1090, 710)),
        ("People and Settings", (510, 590, 780, 710)),
    ]
    for label, box in modules:
        draw_box(draw, box, label, "#F8FAFC", outline="#64748B", font_obj=box_font)
        draw_arrow(draw, ((box[0] + box[2]) // 2, (box[1] + box[3]) // 2), ((center[0] + center[2]) // 2, (center[1] + center[3]) // 2))
    draw.text((90, 800), "Operational records update inventory, production, financial reports, notifications, and user-facing dashboard summaries.", fill="#475569", font=small_font)
    img.save(MODULE_MAP)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.bold = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = BLUE
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = BLUE
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = DARK_BLUE

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Farm Web App Documentation")
    font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Development of a Web-Based Poultry Farm Record and Inventory Management System")
    font(r, size=9, color=MUTED)


def cover_page(doc):
    add_paragraph(doc, before=52, after=10, align=WD_ALIGN_PARAGRAPH.CENTER).add_run("SYSTEM DOCUMENTATION")
    for run in doc.paragraphs[-1].runs:
        font(run, size=12, color=BLUE, bold=True)
    title = (
        "Development of a Web-Based Poultry Farm Record and Inventory Management "
        "System for Chicken Farms in Tambulig, Zamboanga del Sur"
    )
    p = add_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=12)
    r = p.add_run(title)
    font(r, size=23, color=INK, bold=True)
    p = add_paragraph(
        doc,
        "A source-based documentation report prepared from the Django project located at "
        f"{ROOT}",
        before=8,
        after=26,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for run in p.runs:
        font(run, size=11, color=MUTED)
    add_table(
        doc,
        ["Document Item", "Details"],
        [
            ("Application Type", "Web-based poultry farm record and inventory management system"),
            ("Primary Users", "Administrator, farm owner, staff or caretaker"),
            ("Prepared Date", date.today().strftime("%B %d, %Y")),
            ("Reference Format", "Aligned with the supplied thesis-style chapter sample"),
        ],
        [2600, 6500],
        header_fill=LIGHT_BLUE,
    )
    p = add_paragraph(doc, "Abstract", style=None, before=16, after=4)
    for run in p.runs:
        font(run, size=13, color=BLUE, bold=True)
    add_paragraph(
        doc,
        "This documentation describes the design, implementation, and evaluation plan for a Django-based poultry farm system. "
        "The application centralizes flock records, inventory, feeding, medicine and vaccination, egg production, mortality, sales, expenses, reports, notifications, and settings. "
        "The documentation follows the chapter structure of the supplied sample while grounding the content in the actual source code, models, views, templates, and setup files of the farm web app.",
        line=1.10,
    )
    doc.add_page_break()


def chapter_one(doc):
    add_heading(doc, "CHAPTER I", 1)
    add_heading(doc, "BACKGROUND OF THE STUDY", 1)
    add_heading(doc, "Introduction", 2)
    add_paragraph(
        doc,
        "Chicken farms generate daily operational records: flock counts, feed usage, medicine and vaccine schedules, egg production, mortality events, buyers, suppliers, expenses, sales, and reports. "
        "When these records are separated across notebooks or spreadsheets, owners and caretakers can lose visibility into stock levels, treatment schedules, and farm profitability. "
        "The developed farm web app addresses this problem by providing a centralized web-based record and inventory system for chicken farms in Tambulig, Zamboanga del Sur.",
    )
    add_paragraph(
        doc,
        "Poultry production remains an important part of Philippine agriculture, and official livestock and poultry statistics are regularly tracked by the Philippine Statistics Authority [1]. "
        "Good flock management also depends on practical monitoring of production, health, feeding, and environmental factors; veterinary guidance for laying chickens emphasizes management data such as egg production, feed consumption, mortality, and flock condition [2].",
    )
    add_heading(doc, "Project Context", 2)
    add_paragraph(
        doc,
        "The application is designed for a local poultry farm setting where farm owners, staff, and caretakers need a simple but complete system for daily records. "
        "The project implements a Django web application with local static assets, template-driven screens, responsive tables, dashboard cards, charts, modal forms, and Excel exports. "
        "Its source modules show that records are organized around farm operations rather than generic accounting alone.",
    )
    add_heading(doc, "Purpose and Description", 2)
    add_paragraph(
        doc,
        "The purpose of the system is to improve poultry farm recordkeeping by replacing scattered manual records with an integrated web application. "
        "The system allows users to create, update, review, and report on flock, inventory, production, health, sales, expense, buyer, supplier, notification, and user records. "
        "It also supports automatic stock deductions when feed, medicine, or vaccines are used, low-stock and expired-item indicators, vaccination reminders, and summarized production and financial reports.",
    )
    add_heading(doc, "Objectives of the Study", 2)
    add_paragraph(doc, "The general objective is to develop a web-based poultry farm record and inventory management system for chicken farms in Tambulig, Zamboanga del Sur.")
    for item in [
        "Manage user accounts for administrators, farm owners, and staff or caretakers.",
        "Record and monitor chicken flock batches, growing stages, current stock, mortality, and sales-related stock movement.",
        "Track feed, medicine, vaccine, vitamin, equipment, and supply inventory with reorder and expiration alerts.",
        "Record feeding, medicine, vaccination, egg production, mortality, sales, expenses, buyers, and suppliers.",
        "Generate dashboard summaries, reports, Excel exports, notifications, and activity logs for decision support.",
        "Provide a responsive and locally hosted interface using Django templates, Tailwind CSS, HTMX, Alpine.js, and Chart.js.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Scope and Limitations of the Study", 2)
    add_paragraph(
        doc,
        "The scope covers the local Django web app in this workspace. It includes authentication, dashboard analytics, CRUD screens, report pages, Excel exports, local static assets, seed data, and SQLite development storage. "
        "The system is suitable for development, demonstration, and local farm operations after deployment hardening.",
    )
    for item in [
        "The current configuration uses SQLite for development and does not yet include a production PostgreSQL configuration file.",
        "The app does not include online payment processing, SMS delivery, barcode scanning, or native mobile offline synchronization.",
        "Role labels exist in the custom user model, but fine-grained permission enforcement should be expanded before production use.",
        "Some export functions should be regression-tested because report export fields must match the final model field names.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Significance of the Study", 2)
    add_paragraph(doc, "The system benefits the following stakeholders:")
    add_table(
        doc,
        ["Stakeholder", "Benefit"],
        [
            ("Farm owner", "Views flock status, expenses, sales, production, inventory alerts, and profit/loss indicators in one place."),
            ("Caretaker or staff", "Records daily feeding, medicine, vaccination, egg production, mortality, and farm transactions with less repetitive manual work."),
            ("Administrator", "Manages users, activity logs, system settings, reports, and master data such as units and categories."),
            ("Suppliers and buyers", "Benefit indirectly from more organized transaction histories and contact records."),
            ("Future researchers", "Can use the project as a reference for poultry farm management systems using Django and local-first web assets."),
        ],
        [2200, 6900],
    )
    add_heading(doc, "Definition of Terms", 2)
    add_table(
        doc,
        ["Term", "Operational Definition"],
        [
            ("Flock Batch", "A grouped record of chickens with batch number, breed, quantity, age, source, acquisition date, stage, and status."),
            ("Inventory Item", "A feed, medicine, vaccine, vitamin, equipment, or supply item tracked by category, unit, quantity, supplier, cost, expiration date, and reorder level."),
            ("Feeding Record", "A daily or scheduled entry that deducts feed usage from inventory and links the usage to a flock batch."),
            ("Medicine Record", "A record of medicine, vaccine, or vitamin use, including dosage, route, date administered, next schedule, and linked inventory item."),
            ("Mortality Record", "A record of chicken deaths by flock, date, quantity, cause, symptoms, action taken, and remarks."),
            ("Activity Log", "A system-generated trace of important create, update, delete, and view actions for audit visibility."),
        ],
        [2100, 7000],
    )


def chapter_two(doc):
    add_heading(doc, "CHAPTER II", 1)
    add_heading(doc, "REVIEW OF RELATED STUDIES AND SYSTEM", 1)
    add_paragraph(
        doc,
        "This chapter summarizes the technical background, hardware and software requirements, programming environment, and related sources that support the development of the poultry farm web application.",
    )
    add_heading(doc, "Technical Background", 2)
    add_paragraph(
        doc,
        "The project uses Django as its backend framework. Django documentation describes models as the single definitive source of information about stored data, including fields and behavior [3]. "
        "The app follows that pattern through separate Django apps for accounts, dashboard, flocks, inventory, feeding, medicine, eggs, mortality, sales, expenses, suppliers, buyers, reports, notifications, and settings.",
    )
    add_heading(doc, "Hardware Specification", 2)
    add_table(
        doc,
        ["Component", "Minimum Recommendation", "Purpose"],
        [
            ("Client device", "Modern laptop, desktop, tablet, or smartphone with updated browser", "Access dashboard, forms, reports, and responsive tables."),
            ("Processor", "Dual-core processor or better", "Runs browser and local development server during demonstration."),
            ("Memory", "4 GB RAM minimum; 8 GB preferred for development", "Supports Django server, browser, and local build tools."),
            ("Storage", "At least 2 GB free for project, virtual environment, database, static files, and exports", "Stores project files, SQLite database, and generated reports."),
            ("Network", "Local network or internet connection for multi-device access", "Allows staff and owners to access the deployed web app."),
        ],
        [1900, 3200, 4000],
    )
    add_heading(doc, "Software Specification", 2)
    add_table(
        doc,
        ["Software", "Use in the System", "Citation"],
        [
            ("Python and Django", "Backend framework, routing, templates, models, forms, authentication, and admin", "[3], [4], [5]"),
            ("uv", "Python package and environment workflow for installing and running the project", "[6]"),
            ("SQLite", "Development database configured in config/settings.py", "Local project source"),
            ("Django Templates", "Server-rendered HTML screens and reusable partials", "[3]"),
            ("Tailwind CSS", "Utility-first CSS classes and local compiled stylesheet", "[7]"),
            ("HTMX", "Partial-page interactions for modal and CRUD workflows", "[8]"),
            ("Alpine.js", "Small client-side interactions such as toggles, dropdowns, and UI state", "Local static asset"),
            ("Chart.js", "Dashboard visualizations for production and mortality trends", "[9]"),
            ("openpyxl", "Excel export generation in reports/views.py", "Local project dependency"),
        ],
        [2200, 5000, 1900],
    )
    add_heading(doc, "Programming Environment", 2)
    add_paragraph(
        doc,
        "The system is developed as a Django project named config with multiple domain apps. The project can be installed through uv sync, migrated through manage.py migrate, seeded through manage.py seed_data, and run through manage.py runserver. "
        "Static JavaScript and CSS libraries are stored locally under the static directory, supporting the project's offline-friendly frontend requirement after dependencies are installed.",
    )
    add_heading(doc, "Related Literature and Systems", 2)
    add_paragraph(
        doc,
        "Small-scale poultry production guidance from the Food and Agriculture Organization discusses the importance of structured poultry management in smallholder and local farming contexts [10]. "
        "For commercial and semi-commercial flocks, veterinary references emphasize monitoring production, feed use, mortality, and health indicators because those records support timely management decisions [2]. "
        "The system responds to these needs by connecting flock, feeding, medicine, egg production, mortality, and inventory records instead of treating them as isolated forms.",
    )
    add_paragraph(
        doc,
        "From a software perspective, the chosen technologies support a maintainable web application. Django provides built-in conventions for models, forms, authentication, views, and URL dispatching [3], [4], [5]. "
        "HTMX supports AJAX-driven interactions through attributes such as request, target, and swap behavior [8]. Tailwind CSS supports utility-based styling [7], and Chart.js provides charting features for browser-based dashboards [9].",
    )
    add_heading(doc, "Synthesis", 2)
    add_paragraph(
        doc,
        "The reviewed sources support the system's design direction: poultry farms need accurate operational records, and a web-based application can organize those records into searchable, reportable, and actionable information. "
        "The app's source code reflects this synthesis by combining domain-specific farm modules with proven web development tools.",
    )


def chapter_three(doc):
    add_heading(doc, "CHAPTER III", 1)
    add_heading(doc, "DESIGN AND METHODOLOGY", 1)
    add_heading(doc, "Development Model", 2)
    add_paragraph(
        doc,
        "The project can be documented using a waterfall-style sequence similar to the reference sample: requirements analysis, design, development, testing, deployment, and maintenance. "
        "This structure is appropriate for an academic system documentation report because it clearly separates planning, implementation, validation, and future improvement activities.",
    )
    add_heading(doc, "Requirement Analysis", 2)
    add_table(
        doc,
        ["Requirement Area", "Implemented or Represented Feature"],
        [
            ("Authentication", "Custom User model with admin, owner, and staff roles; login, logout, profile, and user management routes."),
            ("Dashboard", "Total chickens, active flocks, low stock, expired items, egg totals, mortality, sales, expenses, recent activity, and charts."),
            ("Farm Operations", "Flocks, inventory, feeding, medicine/vaccination, eggs, mortality, sales, and expenses."),
            ("People Records", "Suppliers and buyers with contact details, notes, active status, and transaction relationships."),
            ("Reports", "Production, mortality, sales, expense, inventory, and profit/loss pages with Excel export functions."),
            ("Notifications", "Low-stock, expired-item, vaccination, mortality, and general notification types with activity logs."),
            ("Settings", "Farm profile, system settings, user management shortcut, and configuration records."),
        ],
        [2700, 6400],
    )
    add_heading(doc, "System Architecture", 2)
    doc.add_picture(str(ARCHITECTURE), width=Inches(6.35))
    add_caption(doc, "Figure 3.1 System Architecture of the Farm Web App")
    add_paragraph(
        doc,
        "The architecture uses the browser as the user entry point, Django templates for rendered screens, Django views and forms for request handling, Django models for data access, SQLite for development storage, and local static assets for UI behavior and charts.",
    )
    add_heading(doc, "Module Map", 2)
    doc.add_picture(str(MODULE_MAP), width=Inches(6.35))
    add_caption(doc, "Figure 3.2 Module Map of the Farm Web App")
    add_heading(doc, "User Roles and Use Case Summary", 2)
    add_table(
        doc,
        ["Role", "Main Use Cases"],
        [
            ("Administrator", "Manage users, settings, farm records, reports, notifications, and activity logs."),
            ("Farm Owner", "View dashboard summaries, monitor flock status, review production, sales, expenses, inventory, and reports."),
            ("Staff/Caretaker", "Enter daily operational records such as feeding, medicine, egg production, mortality, sales, and inventory updates."),
        ],
        [2100, 7000],
    )
    add_heading(doc, "Activity Diagram", 2)
    if ACTIVITY.exists():
        doc.add_picture(str(ACTIVITY), width=Inches(6.3))
        add_caption(doc, "Figure 3.3 Activity Diagram / Flowchart for the Poultry Farm System")
    add_heading(doc, "Database Design", 2)
    add_table(
        doc,
        ["Model", "Primary Purpose", "Important Relationships or Methods"],
        [
            ("User", "Stores account identity, role, phone, address, and profile picture.", "Role helper methods: is_admin_user, is_owner, is_staff_member."),
            ("FarmProfile", "Stores farm name, location, type, contact details, and logo.", "Linked to one user."),
            ("FlockBatch", "Stores batch number, breed, original and current quantity, age, source, stage, and status.", "Used by feeding, medicine, egg, mortality, and sales records."),
            ("InventoryItem", "Stores supplies, feed, medicine, vaccine, costs, quantities, expiration, supplier, and reorder levels.", "Methods: is_low_stock and is_expired."),
            ("InventoryTransaction", "Records stock in, stock out, and adjustment events.", "Save method updates item quantity."),
            ("FeedingRecord", "Records feed used by flock and date/time.", "Save method creates stock-out inventory transaction."),
            ("MedicineRecord", "Records medicine, vaccine, vitamin use and next schedule.", "Optional inventory transaction for used item."),
            ("EggProduction", "Records good, cracked, rejected, and total eggs.", "Save method calculates total eggs."),
            ("MortalityRecord", "Records death quantity, cause, symptoms, action, and remarks.", "Save method adjusts flock current quantity."),
            ("SalesRecord", "Records product type, buyer, quantity, unit price, total, payment status, and balance.", "Save method calculates total amount."),
            ("ExpenseRecord", "Records farm costs by category, amount, date, payment method, and receipt.", "Supports profit/loss report."),
            ("Supplier / Buyer", "Stores contact and transaction-related people records.", "Linked to inventory items or sales records."),
            ("Notification / ActivityLog", "Stores alerts and audit-style activity trail.", "Used by notification context processor and list views."),
            ("SystemSetting", "Stores configurable key-value settings.", "get_setting and set_setting helpers."),
        ],
        [1750, 3650, 3700],
        font_size=8.6,
    )
    add_heading(doc, "Module Design", 2)
    add_table(
        doc,
        ["Module", "Main Screens", "Core Data"],
        [
            ("Dashboard", "Dashboard index, farm profile", "KPI cards, charts, recent activity, notifications."),
            ("Flocks", "Flock list, form, detail, delete modal", "Batch number, breed, quantity, stage, status, stock counts."),
            ("Inventory", "Inventory list, form, detail, transaction, categories, units", "Items, units, suppliers, quantities, cost, expiration, reorder level."),
            ("Feeding", "Feeding list, form, history, delete modal", "Flock, feed item, quantity, date, time, remarks."),
            ("Medicine", "Medicine list, form, vaccination schedule, delete modal", "Flock, medicine type, item, dosage, route, next schedule."),
            ("Eggs", "Egg production list, form, delete modal", "Good, cracked, rejected, and total eggs by flock and date."),
            ("Mortality", "Mortality list, form, delete modal", "Quantity, cause, symptoms, action taken, remarks."),
            ("Sales/Expenses", "Sales records, expense records, categories, forms", "Product sales, buyer, payment status, category, amount, receipt."),
            ("Reports", "Production, mortality, sales, expense, inventory, profit/loss", "Aggregations and Excel exports."),
            ("Notifications", "Notification list, mark read, activity log", "Alerts and user activity."),
        ],
        [1600, 3300, 4200],
        font_size=8.6,
    )
    add_heading(doc, "Implementation Activity Plan", 2)
    add_table(
        doc,
        ["Activity", "Resources Required", "Expected Output", "Remarks"],
        [
            ("Set up Django project", "Python 3.13, uv, Django, editor", "Working project with apps and settings", "Completed in workspace."),
            ("Create models and migrations", "Django ORM, SQLite", "Database schema for farm records", "Initial migrations exist."),
            ("Build templates and static UI", "Django templates, Tailwind, HTMX, Alpine, Chart.js", "Responsive dashboard and CRUD pages", "Local static files used."),
            ("Implement reports and alerts", "Django views, openpyxl, notification utilities", "Reports, exports, and alerts", "Export functions need final regression tests."),
            ("Seed and test data", "seed_data command", "Sample accounts and farm records", "Included for demonstration."),
            ("Deploy and maintain", "Production server, database, backups", "Operational farm system", "Recommended after hardening."),
        ],
        [1900, 2400, 2800, 2000],
        font_size=8.6,
    )


def chapter_four(doc):
    add_heading(doc, "CHAPTER IV", 1)
    add_heading(doc, "DEVELOPMENT, TESTING AND IMPLEMENTATION", 1)
    add_heading(doc, "Development", 2)
    add_paragraph(
        doc,
        "The farm web app is organized into Django apps that each own their models, forms, views, URLs, templates, admin registrations, and migrations. "
        "This modular layout improves maintainability because changes to inventory, reports, sales, or mortality can be handled inside their own application folders while sharing the common project configuration.",
    )
    add_heading(doc, "System Design", 2)
    add_paragraph(
        doc,
        "The user interface uses a dashboard layout with sidebar navigation, record tables, modal forms, confirmation dialogs, filters, alerts, and report pages. "
        "Templates are stored under the templates directory with reusable partials and module-specific pages such as flock_list.html, inventory_list.html, feeding_list.html, medicine_list.html, egg_list.html, mortality_list.html, sales_list.html, expense_list.html, and reports pages.",
    )
    add_heading(doc, "Core Implementation Details", 2)
    for item in [
        "Inventory transactions update item quantities for stock in, stock out, and adjustment actions.",
        "Feeding records automatically create stock-out transactions for the selected feed item.",
        "Medicine and vaccination records can deduct inventory when a linked inventory item is used.",
        "Egg production records calculate total eggs from good, cracked, and rejected egg counts.",
        "Mortality records reduce the current quantity of the associated flock batch.",
        "Sales records calculate total amount from quantity and unit price and expose a balance helper.",
        "Dashboard views aggregate farm data over today, seven-day, and thirty-day windows.",
        "Reports summarize recent production, mortality, sales, expenses, inventory alerts, and profit/loss, with Excel export support.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Testing and Implementation Plan", 2)
    add_table(
        doc,
        ["Test Area", "Test Procedure", "Expected Result"],
        [
            ("Authentication", "Log in using seeded admin, owner, and staff accounts.", "Correct redirect to dashboard and role data available."),
            ("Flock CRUD", "Create, edit, view, and delete a flock batch.", "Records persist and list/detail pages update."),
            ("Inventory CRUD", "Create item, adjust stock, set reorder level and expiration date.", "Inventory list, low-stock, and expired reports reflect the change."),
            ("Feeding deduction", "Create feeding record using a feed item.", "Inventory quantity decreases once through transaction record."),
            ("Medicine deduction", "Create medicine/vaccine record with linked inventory item.", "Inventory quantity decreases once and schedule appears when applicable."),
            ("Egg production", "Record good, cracked, and rejected eggs.", "Total eggs are calculated and production report updates."),
            ("Mortality update", "Record mortality for a flock.", "Flock current quantity decreases and mortality report updates."),
            ("Sales and expenses", "Create sale and expense records.", "Dashboard and profit/loss report update totals."),
            ("Exports", "Download report spreadsheets.", "Excel files contain valid headers and row data matching model fields."),
            ("Notifications", "Run check_alerts command.", "Low-stock, expired-item, and vaccination reminders are generated."),
        ],
        [1800, 4300, 3000],
        font_size=8.6,
    )
    add_heading(doc, "Deployment and Setup", 2)
    for item in [
        "Install dependencies with uv sync.",
        "Apply migrations with uv run python manage.py migrate.",
        "Seed sample data with uv run python manage.py seed_data.",
        "Run the development server with uv run python manage.py runserver.",
        "Open the application at http://127.0.0.1:8000/ and the admin panel at http://127.0.0.1:8000/admin/.",
        "Before production deployment, set DEBUG=False, replace the secret key, restrict ALLOWED_HOSTS, configure a production database, collect static files, and set up backups.",
    ]:
        add_numbered(doc, item)
    add_heading(doc, "User Training", 2)
    add_paragraph(
        doc,
        "Training should focus on common farm routines: logging in, reading dashboard cards, adding flock batches, entering daily feeding and egg production, recording medicine and mortality, reviewing inventory alerts, entering sales and expenses, and exporting reports. "
        "Because the intended users may include non-technical farm staff, training should use sample data and short task-based exercises.",
    )


def chapter_five(doc):
    add_heading(doc, "CHAPTER V", 1)
    add_heading(doc, "RESULTS AND DISCUSSION, CONCLUSION, AND RECOMMENDATION", 1)
    add_heading(doc, "Results and Discussion", 2)
    add_paragraph(
        doc,
        "The source review shows that the system implements the major modules required for poultry farm recordkeeping. "
        "The application stores structured records, presents a dashboard, supports CRUD workflows, creates inventory transactions from farm operations, and generates reports. "
        "Its modular Django structure also provides a maintainable foundation for additional production hardening and feature expansion.",
    )
    add_table(
        doc,
        ["Evaluation Criterion", "Current Status"],
        [
            ("Completeness", "Core modules are present for flocks, inventory, feeding, medicine, eggs, mortality, sales, expenses, people records, reports, notifications, accounts, and settings."),
            ("Usability", "Templates provide dashboard cards, tables, forms, modals, filters, and responsive layout patterns."),
            ("Reliability", "Model save methods automate calculations and stock movements, but automated test coverage should be expanded."),
            ("Maintainability", "Project uses separated Django apps, ModelForms, URL modules, reusable templates, and migrations."),
            ("Deployability", "Development setup is documented; production configuration should be hardened before live use."),
        ],
        [2400, 6700],
    )
    add_heading(doc, "Conclusion", 2)
    add_paragraph(
        doc,
        "The developed farm web app provides a clear and practical basis for a web-based poultry farm record and inventory management system for chicken farms in Tambulig, Zamboanga del Sur. "
        "By centralizing operational records and reports, the system can help owners and caretakers monitor inventory, production, mortality, sales, expenses, and alerts more efficiently than manual recordkeeping.",
    )
    add_heading(doc, "Recommendations", 2)
    for item in [
        "Add comprehensive automated tests for model calculations, stock deduction, report exports, permissions, and notification generation.",
        "Strengthen role-based access control so staff, owners, and administrators have clearly separated permissions.",
        "Prepare production settings for PostgreSQL, environment variables, static file serving, HTTPS, backups, and audit retention.",
        "Add PDF report export, printable receipts, and optional barcode or QR-based inventory tracking.",
        "Consider mobile offline capture for farms with unstable connectivity, then synchronize when the device reconnects.",
        "Review report export field names and add regression tests to prevent spreadsheet download errors.",
    ]:
        add_bullet(doc, item)


def references(doc):
    add_heading(doc, "REFERENCES", 1)
    refs = [
        "Philippine Statistics Authority. Livestock and Poultry Statistics. https://psa.gov.ph/statistics/livestock-poultry",
        "MSD Veterinary Manual. Management of Laying Chickens. https://www.msdvetmanual.com/poultry/nutrition-and-management-poultry/management-of-laying-chickens",
        "Django Software Foundation. Django Documentation 6.0. https://docs.djangoproject.com/en/6.0/",
        "Django Software Foundation. Models. https://docs.djangoproject.com/en/6.0/topics/db/models/",
        "Django Software Foundation. User authentication in Django. https://docs.djangoproject.com/en/6.0/topics/auth/",
        "Astral. uv documentation. https://docs.astral.sh/uv/",
        "Tailwind Labs. Styling with utility classes. https://tailwindcss.com/docs/styling-with-utility-classes",
        "Big Sky Software. htmx documentation. https://htmx.org/docs/",
        "Chart.js Contributors. Chart.js documentation. https://www.chartjs.org/docs/latest/",
        "Food and Agriculture Organization of the United Nations. Small-scale poultry production. https://www.fao.org/4/y5169e/y5169e00.htm",
    ]
    for idx, ref in enumerate(refs, 1):
        add_paragraph(doc, f"[{idx}] {ref}", after=4)


def appendix(doc):
    add_heading(doc, "APPENDIX A: PROJECT STRUCTURE SUMMARY", 1)
    add_table(
        doc,
        ["Folder or File", "Description"],
        [
            ("accounts/", "Custom user model, login/logout views, user forms, and account templates."),
            ("dashboard/", "Dashboard aggregations, farm profile, and dashboard templates."),
            ("flocks/", "Flock batch records and flock management screens."),
            ("inventory/", "Inventory categories, units, items, transactions, and stock workflows."),
            ("feeding/", "Feeding records linked to inventory stock-out transactions."),
            ("medicine/", "Medicine, vaccination, vitamin records, and upcoming schedule view."),
            ("eggs/", "Daily egg production records and calculations."),
            ("mortality/", "Mortality and health monitoring records."),
            ("sales/ and expenses/", "Income and cost record modules."),
            ("suppliers/ and buyers/", "Contact and transaction partner records."),
            ("reports/", "Report pages and Excel exports."),
            ("notifications/", "Notifications, activity logs, context processor, and check_alerts command."),
            ("settings_app/", "System setting and farm configuration screens."),
            ("templates/ and static/", "HTML templates, reusable partials, CSS, JavaScript, charts, and local assets."),
        ],
        [2600, 6500],
        font_size=8.8,
    )
    add_heading(doc, "APPENDIX B: SAMPLE LOGIN CREDENTIALS", 1)
    add_table(
        doc,
        ["Username", "Password", "Role"],
        [
            ("admin", "admin123", "Administrator"),
            ("owner", "owner123", "Farm Owner"),
            ("staff", "staff123", "Staff/Caretaker"),
        ],
        [2400, 2400, 4300],
    )


def build():
    DOCS.mkdir(exist_ok=True)
    generate_diagrams()
    doc = Document()
    configure_document(doc)
    cover_page(doc)
    chapter_one(doc)
    doc.add_page_break()
    chapter_two(doc)
    doc.add_page_break()
    chapter_three(doc)
    doc.add_page_break()
    chapter_four(doc)
    doc.add_page_break()
    chapter_five(doc)
    doc.add_page_break()
    references(doc)
    doc.add_page_break()
    appendix(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
